# -*- coding: utf-8 -*-
"""Отчёт по объёму принятых товаров (на основании документов «Приемка»).

Группировка позиций по номенклатуре; объём = количество в единицах измерения,
суммы — без НДС / НДС / с НДС. Поддерживается фильтр по периоду (дата приёмки)
и поставщику.
"""
from collections import OrderedDict

from app.print_form import _parse_date, nds_to_value


def _to_num(value):
    try:
        return float(str(value).replace(',', '.'))
    except (TypeError, ValueError):
        return 0.0


def build_volume_report(priemki, date_from='', date_to='', postavshik=''):
    """Возвращает dict со строками отчёта и итогами.

    :param priemki: список документов приёмки из БД
    :param date_from/date_to: границы периода по дате документа (дд.мм.гггг)
    :param postavshik: ФИО/наименование поставщика (фильтр, регистронезавис.)
    """
    d_from = _parse_date(date_from) if date_from else None
    d_to = _parse_date(date_to) if date_to else None
    filt_p = (postavshik or '').strip().lower()

    rows = OrderedDict()   # ключ: (краткое наименование/наименование, ед.изм.) -> агрегат
    docs_used = 0
    skipped = 0

    for doc in priemki:
        dt = _parse_date(doc.get('дата'))
        if dt:
            if d_from and dt < d_from:
                skipped += 1
                continue
            if d_to and dt > d_to:
                skipped += 1
                continue
        else:
            skipped += 1
            continue
        p = doc.get('postavshik') or {}
        pname = p.get('kratkoe_naimovanie') or p.get('polnoe_naimovanie') or ''
        if filt_p and filt_p not in pname.lower():
            skipped += 1
            continue
        docs_used += 1
        for poz in doc.get('pozicii', []):
            naim = str(poz.get('naimovanie') or '').strip() or '(без наименования)'
            kratkoe = str(poz.get('kratkoe_naimovanie') or '').strip()
            # для отчёта по объёму используется краткое наименование номенклатуры,
            # при его отсутствии — полное наименование
            show_name = kratkoe or naim
            ed = str(poz.get('ed_izm') or '').strip() or '—'
            kolichestvo = _to_num(poz.get('kolichestvo'))
            tsena = _to_num(poz.get('tsena') or poz.get('cena'))
            bez_nds = round(kolichestvo * tsena, 2)
            koef = nds_to_value(poz.get('stavka_nds_text') or poz.get('stavka_nds'))
            nds_sum = round(bez_nds * koef, 2)
            s_nds = round(bez_nds + nds_sum, 2)
            key = (show_name.lower(), ed.lower())
            if key not in rows:
                rows[key] = {
                    'naimovanie': show_name,
                    'polnoe_naimovanie': naim,
                    'ed_izm': ed,
                    'kod_tovara': poz.get('kod_tovara', ''),
                    'kolichestvo': 0.0,
                    'bez_nds': 0.0,
                    'nds': 0.0,
                    's_nds': 0.0,
                    'dokumenty': set(),
                }
            r = rows[key]
            r['kolichestvo'] += kolichestvo
            r['bez_nds'] += bez_nds
            r['nds'] += nds_sum
            r['s_nds'] += s_nds
            r['dokumenty'].add(doc.get('номер') or '№%s' % doc.get('id'))

    lines = []
    tot = {'kolichestvo': 0.0, 'bez_nds': 0.0, 'nds': 0.0, 's_nds': 0.0}
    for i, r in enumerate(rows.values(), start=1):
        kolichestvo = round(r['kolichestvo'], 3)
        line = {
            'nom': i,
            'kod_tovara': r['kod_tovara'],
            'naimovanie': r['naimovanie'],
            'polnoe_naimovanie': r.get('polnoe_naimovanie', r['naimovanie']),
            'ed_izm': r['ed_izm'],
            'kolichestvo': ('%g' % kolichestvo),
            'bez_nds': '%.2f' % r['bez_nds'],
            'nds': '%.2f' % r['nds'],
            's_nds': '%.2f' % r['s_nds'],
            'dokumenty': ', '.join(sorted(r['dokumenty'])),
        }
        lines.append(line)
        tot['kolichestvo'] += kolichestvo
        tot['bez_nds'] += r['bez_nds']
        tot['nds'] += r['nds']
        tot['s_nds'] += r['s_nds']

    return {
        'lines': lines,
        'docs_used': docs_used,
        'skipped': skipped,
        'period': {'date_from': date_from, 'date_to': date_to,
                   'postavshik': postavshik},
        'totals': {
            'positions': len(lines),
            'bez_nds': '%.2f' % tot['bez_nds'],
            'nds': '%.2f' % tot['nds'],
            's_nds': '%.2f' % tot['s_nds'],
        },
    }


def export_volume_xlsx(report, out_path):
    """Выгрузка отчёта в XLSX (альбомная ориентация A4)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Объем принятых товаров'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4

    bold = Font(bold=True)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    thin = Side(style='thin', color='999999')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = PatternFill('solid', fgColor='EEF3F9')

    per = report['period']
    ws.merge_cells('A1:J1')
    ws['A1'] = 'Отчет по объему принятых товаров'
    ws['A1'].font = Font(bold=True, size=14)
    period_txt = ''
    if per['date_from'] or per['date_to']:
        period_txt = 'Период: с %s по %s' % (per['date_from'] or '—',
                                             per['date_to'] or '—')
    if per['postavshik']:
        period_txt += ('; ' if period_txt else '') + 'Поставщик: %s' % per['postavshik']
    ws.merge_cells('A2:J2')
    ws['A2'] = period_txt or 'Период: весь срок'
    ws['A2'].font = Font(italic=True)

    headers = ['№', 'Код товара', 'Краткое наименование', 'Наименование', 'Ед. изм.',
               'Количество (объем)', 'Сумма без НДС, руб.', 'Сумма НДС, руб.',
               'Сумма с НДС, руб.', 'Документы приемки']
    hr = 4
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=hr, column=c, value=h)
        cell.font = bold
        cell.alignment = center
        cell.border = border
        cell.fill = fill

    r = hr + 1
    for line in report['lines']:
        vals = [line['nom'], line['kod_tovara'], line['naimovanie'],
                line.get('polnoe_naimovanie', ''),
                line['ed_izm'], float(line['kolichestvo']),
                float(line['bez_nds']), float(line['nds']),
                float(line['s_nds']), line['dokumenty']]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = border
            cell.alignment = left if c in (3, 4, 10) else center
            if c in (7, 8, 9):
                cell.number_format = '#,##0.00'
        r += 1

    t = report['totals']
    ws.cell(row=r, column=3, value='ИТОГО').font = bold
    ws.cell(row=r, column=7, value=float(t['bez_nds'])).font = bold
    ws.cell(row=r, column=8, value=float(t['nds'])).font = bold
    ws.cell(row=r, column=9, value=float(t['s_nds'])).font = bold
    for c in range(1, 11):
        cell = ws.cell(row=r, column=c)
        cell.border = border
        cell.fill = fill
        if c in (7, 8, 9):
            cell.number_format = '#,##0.00'

    widths = [5, 14, 26, 40, 10, 14, 16, 14, 16, 28]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = 'A%d' % (hr + 1)
    wb.save(out_path)
