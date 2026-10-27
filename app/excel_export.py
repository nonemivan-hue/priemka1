# -*- coding: utf-8 -*-
"""Выгрузка документов программы в Excel (.xlsx).

Приёмка — таблица позиций с реквизитами, итогами и подписями;
Приказ комиссии — шапка и состав комиссии.
Формат: A4, альбомная ориентация (как печатная форма приёмки).
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.print_form import _parse_date, nds_to_value

THIN = Side(style='thin', color='000000')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill('solid', fgColor='D9E1F2')
TITLE_FONT = Font(bold=True, size=13)
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)


def _setup_sheet(ws, widths):
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _title(ws, row, ncols, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.font = TITLE_FONT
    c.alignment = CENTER
    return row + 1


def _kv(ws, row, label, value, ncols):
    """Строка «метка: значение», растянутая на всю ширину."""
    c = ws.cell(row=row, column=1, value=label)
    c.font = BOLD
    c.border = BORDER
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=ncols)
    v = ws.cell(row=row, column=2, value=value)
    v.alignment = WRAP
    for col in range(2, ncols + 1):
        ws.cell(row=row, column=col).border = BORDER
    return row + 1


def _pos_sums(poz):
    kolichestvo = float(poz.get('kolichestvo') or 0)
    cena = float(poz.get('cena') or 0)
    bez = round(kolichestvo * cena, 2)
    koef = nds_to_value(poz.get('stavka_nds_text') or poz.get('stavka_nds'))
    nds = round(bez * koef, 2)
    return kolichestvo, cena, bez, koef, nds, round(bez + nds, 2)


# --------------------------------------------------------------- Приёмка ----
POZ_HEADERS = ['№', 'Код товара', 'Наименование', 'Страна происх.',
               'Код страны', 'Ед. изм.', 'Код ОКИ', 'Количество', 'Цена, руб.',
               'Сумма без НДС, руб.', 'Ставка НДС', 'Сумма НДС, руб.',
               'Сумма с НДС, руб.']
POZ_WIDTHS = [5, 14, 40, 16, 9, 9, 8, 11, 11, 14, 10, 13, 14]


def export_priemka_xlsx(item, constants, out_path):
    uch = constants.get('uchrezhdenie', {})
    wb = Workbook()
    ws = wb.active
    ws.title = 'Приемка'
    n = len(POZ_HEADERS)
    _setup_sheet(ws, POZ_WIDTHS)

    row = _title(ws, 1, n, 'ДОКУМЕНТ «ПРИЕМКА»')
    row = _title(ws, row, n, '(печатная форма — Приложение к распоряжению '
                             'ректора № 3 от 22.01.2024г., А4 альбомная)')
    row += 1

    post = item.get('postavshik') or {}
    if not isinstance(post, dict):
        post = {'polnoe_naimovanie': str(post)}
    otv = item.get('otvetstvenny') or {}
    dok_osn = ' '.join(filter(None, [item.get('osnovanie_vid'),
                                     '№ ' + str(item.get('osnovanie_nom') or ''),
                                     'от ' + str(item.get('osnovanie_data') or '')]))
    rows_kv = [
        ('Номер документа', item.get('номер') or item.get('nomer') or 'б/н'),
        ('Дата документа', item.get('дата') or item.get('data') or ''),
        ('Учреждение (получатель)', uch.get('учреждение_получатель', '')),
        ('Обособленное подразделение', uch.get('обособленное_подразделение', '')),
        ('Структурное подразделение', uch.get('структурное_подразделение', '')),
        ('Поставщик', post.get('polnoe_naimovanie', '')),
        ('Адрес поставщика', post.get('adres', '')),
        ('ИНН/ОГРН поставщика', '%s / %s' % (post.get('inn', ''), post.get('ogrn', ''))),
        ('Основание (документ-основание)', dok_osn),
        ('Документ об отгрузке', item.get('dok_otgruzka', '')),
        ('ИКЗ', item.get('ikz', '')),
        ('Ответственный за приёмку', '%s, %s' % (otv.get('fio', ''),
                                                 otv.get('dolzhnost', ''))
         if otv else ''),
    ]
    for label, value in rows_kv:
        row = _kv(ws, row, label, value, n)

    row += 1
    # таблица позиций
    hdr_row = row
    for ci, h in enumerate(POZ_HEADERS, start=1):
        c = ws.cell(row=hdr_row, column=ci, value=h)
        c.font = BOLD
        c.fill = HEAD_FILL
        c.border = BORDER
        c.alignment = CENTER
    row += 1

    it_bez = it_nds = it_full = 0.0
    pozicii = item.get('pozicii') or []
    for i, poz in enumerate(pozicii, start=1):
        kolichestvo, cena, bez, _koef, nds, full = _pos_sums(poz)
        it_bez += bez
        it_nds += nds
        it_full += full
        strana = '%s (%s)' % (poz.get('strana_name', ''), poz.get('strana_kod', '')) \
            if poz.get('strana_name') or poz.get('strana_kod') else ''
        vals = [poz.get('nom', i), poz.get('kod_tovara', ''), poz.get('naimovanie', ''),
                poz.get('strana_name', ''), poz.get('strana_kod', ''),
                poz.get('ed_izm', ''), poz.get('oei_kod', ''),
                kolichestvo, cena, bez,
                poz.get('stavka_nds_text') or poz.get('stavka_nds', ''),
                nds, full]
        for ci, v in enumerate(vals, start=1):
            c = ws.cell(row=row, column=ci, value=v)
            c.border = BORDER
            if ci in (9, 10, 12, 13):
                c.number_format = '#,##0.00'
            if ci == 8:
                c.number_format = '#,##0.###'
            if ci == 3:
                c.alignment = WRAP
        row += 1

    if not pozicii:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n)
        c = ws.cell(row=row, column=1, value='Позиции не указаны')
        c.border = BORDER
        row += 1

    # итого
    c = ws.cell(row=row, column=1, value='Итого')
    c.font = BOLD
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    for ci, v in [(10, round(it_bez, 2)), (12, round(it_nds, 2)), (13, round(it_full, 2))]:
        cc = ws.cell(row=row, column=ci, value=v)
        cc.font = BOLD
        cc.number_format = '#,##0.00'
    for ci in range(1, n + 1):
        ws.cell(row=row, column=ci).border = BORDER
    row += 2

    # подписи
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n)
    ruk = uch.get('rukovoditel', '')
    ws.cell(row=row, column=1,
            value='Руководитель _______________ %s          '
                  'Ответственный _______________ %s' %
                  (ruk, otv.get('fio', '') if isinstance(otv, dict) else '')).alignment = WRAP

    wb.save(out_path)
    return out_path


# ----------------------------------------------------------- Приказ комиссии -
PRIKAZ_HEADERS = ['№ п/п', 'Статус в комиссии', 'ФИО', 'Должность']
PRIKAZ_WIDTHS = [8, 24, 45, 45]


def export_prikaz_xlsx(item, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Приказ'
    n = len(PRIKAZ_HEADERS)
    _setup_sheet(ws, PRIKAZ_WIDTHS)

    row = _title(ws, 1, n, 'ПРИКАЗ О СОСТАВЕ КОМИССИИ')
    row += 1
    for label, value in [('Номер приказа', item.get('nomer', '')),
                         ('Дата приказа', item.get('data', '')),
                         ('Наименование комиссии', item.get('nanimovanie', ''))]:
        row = _kv(ws, row, label, value, n)
    row += 1

    for ci, h in enumerate(PRIKAZ_HEADERS, start=1):
        c = ws.cell(row=row, column=ci, value=h)
        c.font = BOLD
        c.fill = HEAD_FILL
        c.border = BORDER
        c.alignment = CENTER
    row += 1

    i = 1
    pred = item.get('predsedatel', '')
    if pred:
        dolzh = item.get('predsedatel_dolzhnost', '')
        vals = [i, 'Председатель комиссии', pred, dolzh]
        for ci, v in enumerate(vals, start=1):
            c = ws.cell(row=row, column=ci, value=v)
            c.border = BORDER
            if ci >= 3:
                c.alignment = WRAP
        row += 1
        i += 1
    for m in item.get('chleny', []) or []:
        if isinstance(m, dict):
            fio, dolzh = m.get('fio', ''), m.get('dolzhnost', '')
        else:
            fio, dolzh = str(m), ''
        vals = [i, 'Член комиссии', fio, dolzh]
        for ci, v in enumerate(vals, start=1):
            c = ws.cell(row=row, column=ci, value=v)
            c.border = BORDER
            if ci >= 3:
                c.alignment = WRAP
        row += 1
        i += 1

    wb.save(out_path)
    return out_path
