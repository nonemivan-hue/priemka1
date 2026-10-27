# -*- coding: utf-8 -*-
"""Отчёт «Количество номенклатуры за период по поставщикам».

На основании документов «Приемка» за выбранный период считается количество
уникальных позиций номенклатуры по каждому поставщику, а также суммарное
количество уникальной номенклатуры (без повторов у разных поставщиков).

Поддерживается разрез отчёта:
  * по документам приёмки (поставщик + номер/дата документа) — детализация;
  * по поставщикам (список наименований номенклатуры через запятую).
"""
from collections import OrderedDict

from app.print_form import _parse_date
from app.report_volume import _to_num


def _poz_name(poz):
    """Наименование позиции номенклатуры для отчёта.

    Приоритет: краткое наименование -> полное наименование -> код товара.
    Для подсчёта уникальной номенклатуры используется код товара (если задан),
    иначе — нормализованное наименование.
    """
    naim = str(poz.get('naimovanie') or '').strip()
    kratkoe = str(poz.get('kratkoe_naimovanie') or '').strip()
    kod = str(poz.get('kod_tovara') or '').strip()
    show = kratkoe or naim or kod or '(без наименования)'
    key = kod.lower() if kod else show.lower()
    return show, key


def build_nomenclature_report(priemki, date_from='', date_to='',
                              postavshik='', group_by='dokument'):
    """Возвращает dict со строками отчёта и итогами.

    :param priemki: список документов приёмки из БД
    :param date_from/date_to: границы периода по дате документа (дд.мм.гггг)
    :param postavshik: фильтр по поставщику (регистронезависимое вхождение)
    :param group_by: 'dokument' — строки по документам приёмки,
                     'postavshik' — одна строка на поставщика
    """
    d_from = _parse_date(date_from) if date_from else None
    d_to = _parse_date(date_to) if date_to else None
    filt_p = (postavshik or '').strip().lower()
    group_by = 'postavshik' if group_by == 'postavshik' else 'dokument'

    rows = OrderedDict()      # ключ строки -> агрегат
    unique_names = {}         # ключ номенклатуры -> отображаемое имя (глобально)
    docs_used = 0
    skipped = 0

    for doc in priemki:
        dt = _parse_date(doc.get('дата'))
        if dt:
            if d_from and dt < d_from:
                skipped += 1
                continue
            if d_to and dt > d_to:
                skipped += 1
                continue
        p = doc.get('postavshik') or {}
        pname = p.get('kratkoe_naimovanie') or p.get('polnoe_naimovanie') or ''
        pname = pname.strip() or '(поставщик не указан)'
        if filt_p and filt_p not in pname.lower():
            skipped += 1
            continue
        docs_used += 1
        nom_keys = []
        for poz in doc.get('pozicii', []):
            show, key = _poz_name(poz)
            unique_names.setdefault(key, show)
            kolichestvo = _to_num(poz.get('kolichestvo'))
            tsena = _to_num(poz.get('tsena') or poz.get('cena'))
            bez_nds = round(kolichestvo * tsena, 2)
            nom_keys.append((show, key, bez_nds))
            if group_by == 'postavshik':
                rkey = pname.lower()
                meta = {'postavshik': pname,
                        'dokument': '',
                        'data_doka': doc.get('дата') or ''}
            else:
                rkey = (pname.lower(), doc.get('номер') or '', doc.get('id'))
                meta = {'postavshik': pname,
                        'dokument': doc.get('номер') or 'б/н',
                        'data_doka': doc.get('дата') or ''}
            r = rows.get(rkey)
            if r is None:
                r = dict(meta)
                r['keys'] = set()
                r['names'] = []
                r['bez_nds'] = 0.0
                rows[rkey] = r
            if key not in r['keys']:
                r['keys'].add(key)
                r['names'].append(show)
            r['bez_nds'] += bez_nds

    lines = []
    tot_count = 0
    tot_sum = 0.0
    for i, r in enumerate(sorted(rows.values(),
                                 key=lambda x: (x['postavshik'].lower(),
                                                x['data_doka'],
                                                x['dokument'].lower())), start=1):
        count = len(r['keys'])
        line = {
            'nom': i,
            'postavshik': r['postavshik'],
            'dokument': r['dokument'],
            'data_doka': r['data_doka'],
            'kolichestvo_nom': count,
            'nomenklatura': ', '.join(r['names']),
            'bez_nds': '%.2f' % r['bez_nds'],
        }
        lines.append(line)
        tot_count += count
        tot_sum += r['bez_nds']

    per = {'date_from': date_from, 'date_to': date_to,
           'postavshik': postavshik, 'group_by': group_by}
    if group_by == 'postavshik':
        period_txt = 'Период: с %s по %s' % (date_from or '—', date_to or '—')
        title = 'Количество номенклатуры за период по поставщикам'
        totals = {
            'postavshikov': len(lines),
            'unikalnoy': len(unique_names),
            'vsego_strok': tot_count,
            'docs_used': docs_used,
            'bez_nds': '%.2f' % tot_sum,
        }
    else:
        period_txt = 'Период: с %s по %s (разрез: по документам приёмки)' % (
            date_from or '—', date_to or '—')
        title = 'Количество номенклатуры за период по документам приёмки'
        totals = {
            'dokumentov': len(lines),
            'unikalnoy': len(unique_names),
            'vsego_strok': tot_count,
            'docs_used': docs_used,
            'bez_nds': '%.2f' % tot_sum,
        }

    return {
        'lines': lines,
        'docs_used': docs_used,
        'skipped': skipped,
        'period': per,
        'title': title,
        'period_txt': period_txt,
        'totals': totals,
    }


def export_nomenclature_xlsx(report, out_path):
    """Выгрузка отчёта о количестве номенклатуры в XLSX (альбомная A4)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Номенклатура по поставщикам'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4

    bold = Font(bold=True)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    thin = Side(style='thin', color='999999')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = PatternFill('solid', fgColor='EEF3F9')

    t = report['totals']
    by_p = report['period']['group_by'] == 'postavshik'

    ws.merge_cells('A1:F1')
    ws['A1'] = report['title']
    ws['A1'].font = Font(bold=True, size=14)
    ws.merge_cells('A2:F2')
    ws['A2'] = report['period_txt']
    ws['A2'].font = Font(italic=True)
    ws.merge_cells('A3:F3')
    if by_p:
        ws['A3'] = ('Поставщиков: %s; всего строк номенклатуры (с повторами): %s; '
                    'уникальной номенклатуры: %s; документов учтено: %s'
                    % (t['postavshikov'], t['vsego_strok'], t['unikalnoy'],
                       t['docs_used']))
    else:
        ws['A3'] = ('Документов приёмки: %s; всего строк номенклатуры (с повторами): %s; '
                    'уникальной номенклатуры: %s'
                    % (t['dokumentov'], t['vsego_strok'], t['unikalnoy']))

    headers = ['№', 'Поставщик']
    if not by_p:
        headers += ['№ документа', 'Дата документа']
    headers += ['Кол-во позиций номенклатуры', 'Наименования номенклатуры',
                'Сумма без НДС, руб.']
    hr = 5
    ncols = len(headers)
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=hr, column=c, value=h)
        cell.font = bold
        cell.alignment = center
        cell.border = border
        cell.fill = fill

    r = hr + 1
    for line in report['lines']:
        vals = [line['nom'], line['postavshik']]
        if not by_p:
            vals += [line['dokument'], line['data_doka']]
        vals += [line['kolichestvo_nom'], line['nomenklatura'],
                 float(line['bez_nds'])]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = border
            col_name = headers[c - 1]
            cell.alignment = left if 'Наименован' in col_name or 'Поставщик' in col_name \
                else center
            if 'Сумма' in col_name:
                cell.number_format = '#,##0.00'
        r += 1

    ws.cell(row=r, column=2, value='ИТОГО').font = bold
    col_count = 3 if by_p else 5
    col_sum = ncols
    ws.cell(row=r, column=col_count, value=t['vsego_strok']).font = bold
    ws.cell(row=r, column=col_sum, value=float(t['bez_nds'])).font = bold
    for c in range(1, ncols + 1):
        cell = ws.cell(row=r, column=c)
        cell.border = border
        cell.fill = fill
        if c == col_sum:
            cell.number_format = '#,##0.00'

    widths = [5, 32] + ([] if by_p else [12, 14]) + [14, 60, 16]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = 'A%d' % (hr + 1)
    wb.save(out_path)
